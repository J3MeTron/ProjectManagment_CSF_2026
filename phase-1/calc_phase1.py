"""Фаза 1: независимый расчет нечеткого коллективного TOPSIS и четкого коллективного МАИ.

Читает экспертные оценки из TOPSIS.xlsx и МАИ.xlsx (те же желтые ячейки, что заполняют эксперты)
и пересчитывает результат на Python. Нужен для сверки с формулами Excel.

Запуск:  python calc_phase1.py [папка_с_книгами]
Зависимости: openpyxl, numpy
"""
import re
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

# Раскладка листов (совпадает с генератором книг)
EXP_ROW0, ALT_ROW0, CRIT_ROW0 = 4, 10, 17
TOPSIS_BLK, AHP_BLK = 13, 13
RI = {1: 0, 2: 0, 3: 0.58, 4: 0.90, 5: 1.12, 6: 1.24, 7: 1.32, 8: 1.41, 9: 1.45}


def num(v):
    """Значение ячейки: число или простая формула вида =1/3."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        m = re.fullmatch(r"=?\s*([\d.,]+)\s*(?:/\s*([\d.,]+))?\s*", v)
        if m:
            a = float(m.group(1).replace(",", "."))
            return a / float(m.group(2).replace(",", ".")) if m.group(2) else a
    raise ValueError(f"не заполнено или не число: {v!r}")


def read_reference(wb):
    ws = wb["Справочник"]
    experts, alts, crits, types = [], [], [], []
    r = EXP_ROW0
    while ws.cell(r, 1).value:
        experts.append(ws.cell(r, 2).value); r += 1
    r = ALT_ROW0
    while ws.cell(r, 1).value:
        alts.append(ws.cell(r, 2).value); r += 1
    r = CRIT_ROW0
    while ws.cell(r, 1).value:
        crits.append(ws.cell(r, 2).value); types.append(ws.cell(r, 3).value); r += 1
    return experts, alts, crits, types


# ------------------------------------------------------------------ TOPSIS
def topsis(path):
    wb = load_workbook(path)
    experts, alts, crits, types = read_reference(wb)
    E, A, C = len(experts), len(alts), len(crits)
    ws = wb["Ввод"]
    w = np.zeros((E, C))
    x = np.zeros((E, A, C))
    for e in range(E):
        r0 = 3 + e * TOPSIS_BLK
        for j in range(C):
            w[e, j] = num(ws.cell(r0 + 2, 2 + j).value)
            for i in range(A):
                x[e, i, j] = num(ws.cell(r0 + 5 + i, 2 + j).value)

    # три упорядоченные оценки -> ТНЧ (пессимистичная, ожидаемая, оптимистичная)
    tfn = lambda arr: np.stack([arr.min(0), np.median(arr, 0), arr.max(0)], -1)
    W = tfn(w) / 9                     # (C, 3)
    X = tfn(x)                         # (A, C, 3)

    N = np.empty_like(X)
    for j in range(C):
        if types[j] == "-":
            N[:, j, :] = X[:, j, 0].min() / X[:, j, ::-1]
        else:
            N[:, j, :] = X[:, j, :] / X[:, j, 2].max()
    V = N * W[None, :, :]
    a_pos = V[:, :, 2].max(0)          # (C,)
    a_neg = V[:, :, 0].min(0)
    d = lambda ref: np.sqrt(((V - ref[None, :, None]) ** 2).mean(-1)).sum(1)
    d_pos, d_neg = d(a_pos), d(a_neg)
    cc = d_neg / (d_pos + d_neg)

    print("=== Нечеткий коллективный TOPSIS ===")
    print("Нечеткие веса критериев (l, m, u) / 9:")
    for j in range(C):
        print(f"  {crits[j]:<40} {W[j].round(3)}")
    print(f"{'Альтернатива':<14}{'d+':>9}{'d-':>9}{'CC':>9}  Ранг")
    order = (-cc).argsort()
    rank = np.empty(A, int); rank[order] = np.arange(1, A + 1)
    for i in range(A):
        print(f"{alts[i]:<14}{d_pos[i]:9.4f}{d_neg[i]:9.4f}{cc[i]:9.4f}  {rank[i]}")
    return dict(zip(alts, cc))


# ------------------------------------------------------------------ МАИ
def priorities(M):
    g = np.prod(M, 1) ** (1 / len(M))
    w = g / g.sum()
    lam = ((M @ w) / w).mean()
    n = len(M)
    ci = (lam - n) / (n - 1)
    cr = ci / RI[n] if RI[n] else 0.0
    return w, lam, cr


def ahp(path):
    wb = load_workbook(path)
    experts, alts, crits, _ = read_reference(wb)
    E, A, C = len(experts), len(alts), len(crits)
    sizes = [C] + [A] * C
    mats = [[None] * E for _ in sizes]
    for e in range(E):
        ws = wb[f"Эксперт{e + 1}"]
        for m, n in enumerate(sizes):
            r0 = 4 + m * AHP_BLK
            M = np.ones((n, n))
            for i in range(n):
                for j in range(i + 1, n):
                    M[i, j] = num(ws.cell(r0 + 2 + i, 2 + j).value)
                    M[j, i] = 1 / M[i, j]
            mats[m][e] = M

    print("\n=== Четкий коллективный МАИ (AIJ, среднее геометрическое суждений) ===")
    names = ["Критерии"] + [f"Альтернативы | {c}" for c in crits]
    print("Индивидуальные CR (контроль согласованности каждого эксперта):")
    for m, nm in enumerate(names):
        crs = [priorities(mats[m][e])[2] for e in range(E)]
        print(f"  {nm:<55} " + "  ".join(f"{c:.3f}" for c in crs))

    agg = [np.exp(np.mean([np.log(M) for M in ms], 0)) for ms in mats]
    res = [priorities(M) for M in agg]
    wc = res[0][0]
    local = np.column_stack([res[1 + j][0] for j in range(C)])  # (A, C)
    glob = local @ wc
    print("Коллективные матрицы: lambda max / CR")
    for m, nm in enumerate(names):
        print(f"  {nm:<55} {res[m][1]:.3f} / {res[m][2]:.3f}")
    print("Веса критериев:")
    for j in range(C):
        print(f"  {crits[j]:<40} {wc[j]:.4f}")
    order = (-glob).argsort()
    rank = np.empty(A, int); rank[order] = np.arange(1, A + 1)
    print(f"{'Альтернатива':<14}{'Приоритет':>11}  Ранг")
    for i in range(A):
        print(f"{alts[i]:<14}{glob[i]:11.4f}  {rank[i]}")
    return dict(zip(alts, glob))


if __name__ == "__main__":
    folder = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent)
    topsis(folder / "TOPSIS.xlsx")
    ahp(folder / "МАИ.xlsx")
